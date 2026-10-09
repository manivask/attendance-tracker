import openpyxl
import json
import urllib.request
from datetime import datetime

WEBHOOK_URL = "https://script.google.com/macros/s/AKfycbw_v-oyFIx4VWVsdkfTL7x0iykVJ34uK6HnzBZP0bxyenYBWX301vkN2c97ixSBDlW8ZA/exec"

def parse_date_str(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s:
        return None
    # Try parsing common formats
    for fmt in ["%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%m-%d-%Y", "%d-%b-%Y"]:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None

def main():
    print("Loading TBTA-2026-2027- RHS-Student_Attendance.xlsx...")
    wb = openpyxl.load_workbook("TBTA-2026-2027- RHS-Student_Attendance.xlsx", data_only=True)

    # First fetch current cloud state so we don't lose any existing test/homework marks
    print("Fetching current cloud state from Google Sheets...")
    req = urllib.request.Request(WEBHOOK_URL, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            data = json.loads(resp.read().decode())
            current_cloud_state = data.get("state") or {}
            print("Successfully fetched cloud state.")
    except Exception as e:
        print(f"Warning: could not fetch existing cloud state: {e}")
        current_cloud_state = {}

    attendance = current_cloud_state.get("attendance") or {"Teachers": {}, "Students": {}}
    if "Teachers" not in attendance: attendance["Teachers"] = {}
    if "Students" not in attendance: attendance["Students"] = {}

    homework = current_cloud_state.get("homework") or {}
    tests = current_cloud_state.get("tests") or {}
    locked_dates = current_cloud_state.get("lockedDates") or []
    attestations = current_cloud_state.get("attestations") or {}

    class_sheets = ["illanthalir", "Mazhalai", "Nilai-1", "Nilai-2", "Nilai-3", "Nilai-4", "Nilai-5", "Nilai-6", "Nilai-7", "Nilai-8"]
    
    # Load students and teachers
    students_list = []
    teachers_list = []
    student_name_to_id = {}
    teacher_name_to_id = {}

    # If students already in cloud state, build mapping
    for s in current_cloud_state.get("students", []):
        student_name_to_id[s["Name"].strip()] = s["ID"]
        students_list.append(s)

    for t in current_cloud_state.get("teachers", []):
        teacher_name_to_id[t["Name"].strip()] = t["ID"]
        teachers_list.append(t)

    # Parse Teacher sheet
    if "Teacher" in wb.sheetnames:
        ws_t = wb["Teacher"]
        date_cols = {}
        for c in range(1, ws_t.max_column + 1):
            h_val = ws_t.cell(row=1, column=c).value
            ds = parse_date_str(h_val)
            if ds:
                date_cols[c] = ds

        for r in range(2, ws_t.max_row + 1):
            f_name = ws_t.cell(row=r, column=1).value or ""
            l_name = ws_t.cell(row=r, column=2).value or ""
            cls_assign = ws_t.cell(row=r, column=3).value or ""
            full_name = f"{f_name} {l_name}".strip()
            if not full_name:
                continue

            t_id = teacher_name_to_id.get(full_name, f"T{len(teachers_list)+1:03d}")
            if full_name not in teacher_name_to_id:
                teacher_name_to_id[full_name] = t_id
                teachers_list.append({
                    "ID": t_id,
                    "Name": full_name,
                    "Class Assignment": str(cls_assign),
                    "Location": "Riverview"
                })

            for col_idx, date_str in date_cols.items():
                cell_val = ws_t.cell(row=r, column=col_idx).value
                if cell_val is not None:
                    st_val = str(cell_val).strip().upper()
                    st = "present" if st_val == "P" else "late" if st_val == "L" else "absent" if st_val == "A" else ""
                    if st:
                        if date_str not in attendance["Teachers"]:
                            attendance["Teachers"][date_str] = {}
                        attendance["Teachers"][date_str][t_id] = st

    # Parse each Class sheet for students and attendance
    total_parsed_marks = 0
    for s_name in class_sheets:
        if s_name not in wb.sheetnames:
            continue
        ws = wb[s_name]
        grade_name = s_name.replace("Nilai-", "Nilai ")
        if grade_name == "illanthalir": grade_name = "Ilanthalir"

        date_cols = {}
        for c in range(1, ws.max_column + 1):
            h_val = ws.cell(row=1, column=c).value
            ds = parse_date_str(h_val)
            if ds:
                date_cols[c] = ds

        for r in range(2, ws.max_row + 1):
            cell_1 = ws.cell(row=r, column=1).value or ""
            cell_2 = ws.cell(row=r, column=2).value or ""
            cell_3 = ws.cell(row=r, column=3).value or ""
            
            # Check if this row is student ID or First name
            s_id = ""
            full_name = ""
            if str(cell_1).strip().upper().startswith("S") and len(str(cell_1).strip()) <= 6:
                s_id = str(cell_1).strip()
                full_name = f"{cell_2} {cell_3}".strip()
            else:
                full_name = f"{cell_1} {cell_2}".strip()

            if not full_name or full_name.lower() == "first name":
                continue

            if not s_id:
                s_id = student_name_to_id.get(full_name, f"S{len(students_list)+1:03d}")

            if full_name not in student_name_to_id:
                student_name_to_id[full_name] = s_id
                students_list.append({
                    "ID": s_id,
                    "Name": full_name,
                    "Grade": grade_name,
                    "Location": "Riverview"
                })

            for col_idx, date_str in date_cols.items():
                cell_val = ws.cell(row=r, column=col_idx).value
                if cell_val is not None:
                    st_val = str(cell_val).strip().upper()
                    st = "present" if st_val == "P" else "late" if st_val == "L" else "absent" if st_val == "A" else ""
                    if st:
                        if date_str not in attendance["Students"]:
                            attendance["Students"][date_str] = {}
                        attendance["Students"][date_str][s_id] = st
                        total_parsed_marks += 1

    print(f"Extracted attendance for {len(students_list)} students and {len(teachers_list)} teachers across all recorded dates.")
    print(f"Recorded student attendance dates: {sorted(list(attendance['Students'].keys()))}")
    print(f"Recorded teacher attendance dates: {sorted(list(attendance['Teachers'].keys()))}")
    print(f"Total attendance mark cells processed: {total_parsed_marks}")

    # Build payload
    payload = {
        "version": "1.0",
        "activeDate": "2026-10-02",
        "updatedAt": datetime.utcnow().isoformat() + "Z",
        "updatedBy": "Administrator (Master Copy Sync)",
        "attendance": attendance,
        "homework": homework,
        "tests": tests,
        "lockedDates": locked_dates,
        "attestations": attestations,
        "students": students_list,
        "teachers": teachers_list
    }

    # Push to Google Apps Script Webhook
    print(f"Syncing master copy to Google Sheets & Google Drive Live Sync Webhook...")
    post_data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(WEBHOOK_URL, data=post_data, headers={"Content-Type": "text/plain;charset=utf-8"})
    
    with urllib.request.urlopen(req, timeout=30) as resp:
        result = resp.read().decode("utf-8")
        print(f"Server Response: {result}")
        res_json = json.loads(result)
        if res_json.get("status") == "success":
            print(" Master copy successfully synchronized to Google Drive & Sheets Live Sync!")
        else:
            print(f"❌ Server returned error: {res_json}")

if __name__ == "__main__":
    main()
