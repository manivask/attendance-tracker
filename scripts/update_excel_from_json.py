"""
TBTA Attendance Spreadsheet Synchronizer
Reads data/attendance-state.json and writes attendance records into TBTA-2026-2027- RHS-Student_Attendance.xlsx
"""
import os
import json
import openpyxl

JSON_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "attendance-state.json")
EXCEL_PATH = os.path.join(os.path.dirname(__file__), "..", "TBTA-2026-2027- RHS-Student_Attendance.xlsx")

def sync_json_to_excel():
    if not os.path.exists(JSON_PATH):
        print(f"JSON state file not found at {JSON_PATH}")
        return

    if not os.path.exists(EXCEL_PATH):
        print(f"Excel file not found at {EXCEL_PATH}")
        return

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        state = json.load(f)

    attendance = state.get("attendance", {})
    student_att = attendance.get("Students", {})
    teacher_att = attendance.get("Teachers", {})

    wb = openpyxl.load_workbook(EXCEL_PATH)

    # Process all sheets in workbook
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1] if cell.value is not None]
        if not headers:
            continue

        id_col_idx = 1
        # Find date columns
        date_cols = {}
        for col_idx, h in enumerate(headers, 1):
            h_str = str(h).strip()
            if len(h_str) == 10 and h_str.count("-") == 2:
                date_cols[h_str] = col_idx

        if not date_cols:
            continue

        is_teacher_sheet = sheet_name.lower() in ["teacher", "teachers"]

        for row_idx in range(2, ws.max_row + 1):
            row_id = str(ws.cell(row=row_idx, column=id_col_idx).value or "").strip()
            if not row_id:
                continue

            for date_str, col_idx in date_cols.items():
                if is_teacher_sheet:
                    status = teacher_att.get(date_str, {}).get(row_id)
                else:
                    status = student_att.get(date_str, {}).get(row_id)

                if status == "present":
                    ws.cell(row=row_idx, column=col_idx).value = "P"
                elif status == "absent":
                    ws.cell(row=row_idx, column=col_idx).value = "A"

    wb.save(EXCEL_PATH)
    print(f"Successfully synchronized JSON state to {EXCEL_PATH}")

if __name__ == "__main__":
    sync_json_to_excel()
